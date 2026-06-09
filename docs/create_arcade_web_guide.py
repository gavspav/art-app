from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "docs/Arcade_Cabinet_Web_App_Guide.docx"
BLUE = "145C74"
TEAL = "2F899E"
PALE_BLUE = "DCEFF4"
PALE_GREEN = "E2F2E8"
PALE_ORANGE = "FFF0D7"
PALE_GREY = "ECEFF1"
MID_GREY = "65757C"
WHITE = "FFFFFF"
BLACK = "172126"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=100, bottom=100, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TEAL)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(color)
    run_properties.append(underline)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(run_properties)
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def style_run(run, size=10, bold=False, color=BLACK, font="Arial"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_title(doc, text, subtitle=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    style_run(run, size=24, bold=True, color=BLUE)
    if subtitle:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(12)
        run2 = p2.add_run(subtitle)
        style_run(run2, size=11, color=MID_GREY)


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    style_run(run, size=14, bold=True, color=BLUE)
    return p


def add_note(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(9.6)
    shade(cell, PALE_ORANGE)
    set_cell_margins(cell, 130, 160, 130, 160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    first = p.add_run(f"{label}: ")
    style_run(first, size=9.5, bold=True, color=BLACK)
    second = p.add_run(text)
    style_run(second, size=9.5, color=BLACK)


def add_startup_steps(doc):
    steps = [
        ("Open the app", "Go to ", "https://arcadecabinet--gavxflx.netlify.app/"),
        ("Open Sound", "Press 6 to open the Sound tab.", None),
        ("Enable controls", "Make sure Cab keys and Button pairs are checked.", None),
        ("Start sound", "Click the button to start the audio.", None),
        ("Focus the artwork", "Click once on the canvas.", None),
        ("Go fullscreen", "Press F.", None),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    widths = [Inches(0.55), Inches(1.8), Inches(7.25)]
    headers = ["Step", "Action", "What to do"]
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = widths[index]
        shade(cell, BLUE)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
        style_run(p.add_run(header), size=9, bold=True, color=WHITE)
    set_repeat_table_header(table.rows[0])
    for number, (action, instruction, url) in enumerate(steps, start=1):
        cells = table.add_row().cells
        for index, cell in enumerate(cells):
            cell.width = widths[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 95, 115, 95, 115)
            shade(cell, WHITE if number % 2 else PALE_BLUE)
        p0 = cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(p0.add_run(str(number)), size=10, bold=True, color=BLUE)
        style_run(cells[1].paragraphs[0].add_run(action), size=9.5, bold=True)
        p2 = cells[2].paragraphs[0]
        style_run(p2.add_run(instruction), size=9.5)
        if url:
            add_hyperlink(p2, url, url)


def add_key(cell, key, action, group="control"):
    fills = {
        "system": PALE_ORANGE,
        "joystick": PALE_GREEN,
        "button": PALE_BLUE,
        "unused": PALE_GREY,
    }
    shade(cell, fills[group])
    set_cell_margins(cell, 90, 75, 90, 75)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    key_run = p.add_run(key)
    style_run(key_run, size=11, bold=True, color=BLUE if group != "unused" else MID_GREY)
    if action:
        action_run = p.add_run(f"\n{action}")
        style_run(action_run, size=7.5, bold=True, color=BLACK)


def add_keyboard(doc):
    keys = {
        "6": ("SOUND TAB", "system"),
        "W": ("SIZE +", "joystick"),
        "E": ("BACKGROUND +", "button"),
        "R": ("PALETTE +", "button"),
        "T": ("BLEND TOGGLE", "button"),
        "U": ("WOBBLE +\nNOISE +", "joystick"),
        "I": ("SIDES +", "button"),
        "O": ("OPACITY +", "button"),
        "P": ("CURVINESS\nTOGGLE", "button"),
        "A": ("SPEED -", "joystick"),
        "S": ("SPEED +", "joystick"),
        "F": ("FULLSCREEN", "system"),
        "H": ("LAYERS -", "joystick"),
        "J": ("LAYERS +", "joystick"),
        "Z": ("SIZE -", "joystick"),
        "C": ("BACKGROUND -", "button"),
        "V": ("PALETTE -", "button"),
        "N": ("WOBBLE -\nNOISE -", "joystick"),
        "M": ("SIDES -", "button"),
        ",": ("OPACITY -", "button"),
    }
    rows = [
        ["1", "2", "3", "4", "5", "6", "7", "8", "9", "0", "-"],
        ["Q", "W", "E", "R", "T", "Y", "U", "I", "O", "P", "["],
        ["", "A", "S", "D", "F", "G", "H", "J", "K", "L", ""],
        ["", "Z", "X", "C", "V", "B", "N", "M", ",", ".", ""],
    ]
    table = doc.add_table(rows=0, cols=11)
    table.autofit = False
    for row_keys in rows:
        cells = table.add_row().cells
        for index, key in enumerate(row_keys):
            cells[index].width = Inches(0.88)
            if not key:
                shade(cells[index], WHITE)
                cells[index]._tc.get_or_add_tcPr().append(OxmlElement("w:tcBorders"))
                continue
            action, group = keys.get(key, ("", "unused"))
            add_key(cells[index], key, action, group)


def add_control_summary(doc):
    data = [
        ("Joystick 1", "A / S", "Decrease / increase global speed"),
        ("Joystick 1", "W / Z", "Increase / decrease shape size"),
        ("Joystick 2", "H / J", "Decrease / increase number of layers"),
        ("Joystick 2", "U / N", "Increase / decrease wobble and noise together"),
        ("Button pair", "E / C", "Cycle background colours forward / backward"),
        ("Button pair", "R / V", "Cycle palettes forward / backward"),
        ("Button", "T", "Toggle blend mode between source-over and difference"),
        ("Button pair", "I / M", "Increase / decrease number of sides"),
        ("Button pair", "O / ,", "Increase / decrease opacity"),
        ("Button", "P", "Toggle curviness between 0 and 1"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    widths = [Inches(1.45), Inches(1.1), Inches(7.05)]
    for index, header in enumerate(("Control", "Keys", "Effect")):
        cell = table.rows[0].cells[index]
        cell.width = widths[index]
        shade(cell, BLUE)
        set_cell_margins(cell)
        style_run(cell.paragraphs[0].add_run(header), size=8.5, bold=True, color=WHITE)
    for row_number, row in enumerate(data):
        cells = table.add_row().cells
        for index, text in enumerate(row):
            cell = cells[index]
            cell.width = widths[index]
            set_cell_margins(cell, 45, 100, 45, 100)
            shade(cell, WHITE if row_number % 2 == 0 else PALE_BLUE)
            style_run(cell.paragraphs[0].add_run(text), size=8, bold=index == 1)


def main():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.1

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_run(header.add_run("ARCADE CABINET | OPERATOR GUIDE"), size=8, bold=True, color=MID_GREY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(footer.add_run("Arcade Cabinet Web App"), size=8, color=MID_GREY)

    add_title(doc, "Arcade Cabinet Web App", "Quick-start and keyboard control reference")
    add_heading(doc, "Start The Experience")
    add_startup_steps(doc)
    add_note(
        doc,
        "Audio",
        "Browsers require a click before sound can start. If sound is silent, return to the Sound tab with 6 and click the start-audio button again.",
    )

    doc.add_page_break()
    add_title(doc, "Keyboard Control Layout", "Cab keys and Button pairs must both be checked")
    add_keyboard(doc)
    add_note(
        doc,
        "Hold joysticks",
        "Hold the green joystick keys to continuously change their values. Tap the blue button keys to step, cycle, or toggle their controls.",
    )
    add_heading(doc, "Control Summary")
    add_control_summary(doc)

    doc.core_properties.title = "Arcade Cabinet Web App Operator Guide"
    doc.core_properties.subject = "Startup and keyboard controls"
    doc.core_properties.author = "gavXflx"
    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
