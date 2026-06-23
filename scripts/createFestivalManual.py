from pathlib import Path
import json
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/FESTIVAL_FIELD_MANUAL.md'
PRESET = ROOT / 'src/config/defaultArcadePreset.json'
OUTPUT = ROOT / 'docs/Arcade_Cabinet_Festival_Field_Manual.docx'

NAVY = '18344F'
BLUE = '247BA0'
PALE = 'EAF3F7'
GREY = 'F1F3F5'
GOLD = 'F7E5A1'
WHITE = 'FFFFFF'
DARK = '20262E'


def set_shading(element, color):
    props = element.get_or_add_tcPr() if hasattr(element, 'get_or_add_tcPr') else element.get_or_add_pPr()
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color)
    props.append(shading)


def set_cell(cell, value, bold=False, color=DARK):
    cell.text = ''
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(value))
    run.bold = bold
    run.font.name = 'Calibri'
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document, rows):
    table = document.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for column, value in enumerate(rows[0]):
        set_cell(table.rows[0].cells[column], value, True, WHITE)
        set_shading(table.rows[0].cells[column]._tc, NAVY)
    header_props = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement('w:tblHeader')
    repeat.set(qn('w:val'), 'true')
    header_props.append(repeat)
    for row_number, values in enumerate(rows[1:]):
        cells = table.add_row().cells
        for column, value in enumerate(values):
            set_cell(cells[column], value)
            if row_number % 2:
                set_shading(cells[column]._tc, GREY)
    for row in table.rows:
        no_split = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(no_split)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_inline(paragraph, text):
    pieces = re.split(r'(`[^`]+`|\*\*[^*]+\*\*)', text)
    for piece in pieces:
        if piece.startswith('`') and piece.endswith('`'):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = 'Menlo'
            run.font.size = Pt(8)
        elif piece.startswith('**') and piece.endswith('**'):
            paragraph.add_run(piece[2:-2]).bold = True
        else:
            paragraph.add_run(piece)


def parse_markdown(document, text):
    lines = text.splitlines()
    index = 0
    in_code = False
    while index < len(lines):
        line = lines[index]
        if line.startswith('```'):
            in_code = not in_code
            index += 1
            continue
        if in_code:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.15)
            paragraph.paragraph_format.right_indent = Inches(0.15)
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(line or ' ')
            run.font.name = 'Menlo'
            run.font.size = Pt(7.6)
            set_shading(paragraph._p, GREY)
        elif line.startswith('|') and index + 1 < len(lines) and re.match(r'^\|[\s:|-]+\|$', lines[index + 1]):
            rows = []
            while index < len(lines) and lines[index].startswith('|'):
                values = [part.strip() for part in lines[index].strip('|').split('|')]
                if not all(re.fullmatch(r'[-: ]+', value) for value in values):
                    rows.append(values)
                index += 1
            add_table(document, rows)
            continue
        elif line.startswith('# '):
            document.add_heading(line[2:], 1)
        elif line.startswith('## '):
            document.add_heading(line[3:], 2)
        elif line.startswith('### '):
            document.add_heading(line[4:], 3)
        elif line.startswith('> '):
            table = document.add_table(rows=1, cols=1)
            cell = table.cell(0, 0)
            set_shading(cell._tc, GOLD if 'Important' in line or 'Do not' in line else PALE)
            cell.text = ''
            add_inline(cell.paragraphs[0], line[2:])
        elif re.match(r'^\d+\. ', line):
            paragraph = document.add_paragraph(style='List Number')
            add_inline(paragraph, re.sub(r'^\d+\. ', '', line))
        elif line.startswith('- '):
            paragraph = document.add_paragraph(style='List Bullet')
            add_inline(paragraph, line[2:])
        elif line.strip():
            paragraph = document.add_paragraph()
            add_inline(paragraph, line)
        index += 1


def append_preset_inventory(document):
    preset = json.loads(PRESET.read_text())
    document.add_heading('Appendix: Default Preset Inventory', 1)
    document.add_paragraph('Generated directly from the current defaultArcadePreset.json so this appendix matches the build source.')
    document.add_heading('Top-level keys and counts', 2)
    rows = [['Key', 'Stored form / count']]
    for key, value in preset.items():
        count = f'{len(value)} entries' if isinstance(value, (dict, list)) else type(value).__name__
        rows.append([key, count])
    add_table(document, rows)
    document.add_heading('Parameter keys', 2)
    parameters = preset.get('parameters', [])
    rows = [['Parameter', 'Stored value or summary']]
    parameter_items = (
        sorted(parameters.items())
        if isinstance(parameters, dict)
        else [(value.get('id', str(index)), value) for index, value in enumerate(parameters)]
    )
    for key, value in parameter_items:
        rendered = json.dumps(value, separators=(',', ':'))
        rows.append([key, rendered[:130] + ('...' if len(rendered) > 130 else '')])
    add_table(document, rows)
    document.add_heading('MIDI mappings', 2)
    mappings = preset.get('midiMappings', {})
    items = mappings.items() if isinstance(mappings, dict) else enumerate(mappings)
    rows = [['Mapping key', 'Stored mapping']]
    for key, value in items:
        rendered = json.dumps(value, separators=(',', ':'))
        rows.append([str(key), rendered[:150] + ('...' if len(rendered) > 150 else '')])
    add_table(document, rows)


document = Document()
section = document.sections[0]
section.page_height = Inches(11)
section.page_width = Inches(8.5)
section.top_margin = section.bottom_margin = Inches(0.7)
section.left_margin = section.right_margin = Inches(0.78)
section.header_distance = section.footer_distance = Inches(0.3)

normal = document.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(9.2)
normal.paragraph_format.space_after = Pt(4)
for name, size, color in [('Title', 32, NAVY), ('Heading 1', 21, NAVY), ('Heading 2', 14, BLUE), ('Heading 3', 11, DARK)]:
    style = document.styles[name]
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(9)
    style.paragraph_format.space_after = Pt(4)
document.styles['Heading 1'].paragraph_format.page_break_before = True

header = section.header.paragraphs[0]
header.text = 'GAVXFLX ARCADE CABINET  /  FESTIVAL FIELD MANUAL'
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor.from_string(BLUE)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Page ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

parse_markdown(document, SOURCE.read_text())
append_preset_inventory(document)
document.core_properties.title = 'Arcade Cabinet Festival Field Manual'
document.core_properties.subject = 'Offline operation, troubleshooting, configuration and codebase guide'
document.core_properties.author = 'GAVXFLX / Codex'
document.save(OUTPUT)
print(OUTPUT)
